# importador.py
# Lê listas de insumos de arquivos externos (Excel, CSV, Word, PDF) e
# transforma em linhas com nome, categoria, unidade e quantidade.
import csv
import io
import re
import unicodedata

from insumos_iniciais import UNIDADES

EXTENSOES = {"xlsx", "xlsm", "csv", "txt", "docx", "pdf"}

# Palavras que identificam cada coluna no cabeçalho
CABECALHOS = {
    "nome": ["insumo", "nome", "produto", "item", "descricao", "mercadoria", "material"],
    "categoria": ["categoria", "grupo", "setor", "tipo", "secao"],
    "unidade": ["unidade", "und", "unid", "un", "medida", "um"],
    "quantidade": ["quantidade", "qtd", "qtde", "quant", "contada", "contado", "contagem", "fisico", "peso", "estoque",
                   "saldo"],
    "observacao": ["observacoes", "observacao", "obs", "validade"],
}

# Sinônimos de unidades -> unidade do app
SINONIMOS_UNIDADE = {
    "un": "un", "und": "un", "unid": "un", "unidade": "un", "unidades": "un", "uni": "un", "pc": "un", "pç": "un",
    "peca": "un", "pecas": "un",
    "kg": "kg", "kgs": "kg", "quilo": "kg", "quilos": "kg", "kilo": "kg", "kilos": "kg",
    "g": "g", "gr": "g", "grama": "g", "gramas": "g",
    "l": "L", "lts": "L", "litro": "L", "litros": "L",
    "ml": "ml", "cx": "cx", "caixa": "cx", "caixas": "cx", "pct": "pct", "pacote": "pct", "pacotes": "pct",
    "fd": "fd", "fardo": "fd", "fardos": "fd", "gf": "gf", "garrafa": "gf", "garrafas": "gf",
    "lata": "lt", "latas": "lt", "mc": "mç", "maco": "mç", "macos": "mç", "dz": "dz", "duzia": "dz", "duzias": "dz",
}


def sem_acento(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().lower()


def chave_nome(nome):
    """Nome sem acento e no singular, para casar 'Tomates' com 'Tomate' e 'Camarões' com 'Camarão'."""
    def singular(palavra):
        if len(palavra) <= 3:
            return palavra
        if palavra.endswith(("oes", "aes", "aos")):
            return palavra[:-3] + "ao"
        if palavra.endswith(("res", "zes", "ses")):
            return palavra[:-2]
        return palavra[:-1] if palavra.endswith("s") else palavra
    return " ".join(singular(p) for p in sem_acento(nome).split(" "))


def normalizar_unidade(texto):
    chave = sem_acento(texto).strip(" .")
    if chave in SINONIMOS_UNIDADE:
        return SINONIMOS_UNIDADE[chave]
    for u in UNIDADES:
        if sem_acento(u) == chave:
            return u
    return None


# Abreviações de unidade que aparecem coladas ao nome ("Bife do Vazio Empanado kg", "Arroz (kg)").
# Palavras inteiras como "Lata" ficam no nome: "Coca Cola Lata" é o nome do produto.
SIGLAS_UNIDADE = {"kg", "kgs", "g", "gr", "grs", "mg", "l", "lt", "lts", "ml", "un", "und", "unid", "cx", "pct",
                  "pc", "dz", "fd", "gf", "mc"}


def _unidade_antes(texto):
    """'kg 4,5' / 'KG: 4,5' -> ('4,5', 'kg'); senão (texto, None)."""
    m = re.fullmatch(r"([^\W\d_]+)\.?\s*[:\-]?\s*(\d[\d.,\s]*)", texto)
    if m and normalizar_unidade(m.group(1)):
        return m.group(2).strip(), normalizar_unidade(m.group(1))
    return texto, None


def tirar_unidade_do_nome(nome):
    """'Bife do Vazio Empanado kg' / 'Arroz (kg)' / 'Feijão - KG' -> ('Bife do Vazio Empanado', 'kg')."""
    m = re.fullmatch(r"(.+?)[\s\-–:,]*[(\[]?\s*([^\W\d_]+)\.?\s*[)\]]?", nome.strip())
    if m and m.group(1).strip() and sem_acento(m.group(2)) in SIGLAS_UNIDADE and normalizar_unidade(m.group(2)) \
            and (m.group(1) != nome.strip()):
        resto = m.group(1).strip(" -–:,")
        if resto and re.search(r"[^\W\d_]", resto):
            return resto, normalizar_unidade(m.group(2))
    return nome, None


def separar_quantidade(texto):
    """'4,5 kg' -> (4.5, 'kg'); '4,5' -> (4.5, None); '4,5KG.' -> (4.5, 'kg'). Número com a unidade escrita
    junto vale o mesmo que o número sozinho. (None, None) se não for número."""
    if isinstance(texto, (int, float)):
        return parse_numero(texto), None
    texto = str(texto or "").strip()
    texto, antes = _unidade_antes(texto)
    if antes:
        return parse_numero(texto), antes
    m = re.fullmatch(r"([\d][\d.,\s]*?)\s*([^\W\d_]+)\.?", texto)
    if m:
        unidade = normalizar_unidade(m.group(2))
        if unidade:
            return parse_numero(m.group(1)), unidade
    return parse_numero(texto), None


def parse_numero(texto):
    """'1,5' / '1.5' / '1.250,5' / '1,5 kg' -> float. None se não for número."""
    if isinstance(texto, (int, float)):
        return float(texto) if texto >= 0 else None
    texto = str(texto or "").strip()
    if not texto:
        return None
    texto = _unidade_antes(texto)[0]
    m = re.fullmatch(r"([\d][\d.,\s]*?)\s*([^\W\d_]+)\.?", texto)
    if m and normalizar_unidade(m.group(2)):
        texto = m.group(1).strip()
    texto = texto.replace(" ", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", texto):
        texto = texto.replace(".", "")
    texto = texto.replace(",", ".")
    try:
        valor = float(texto)
    except ValueError:
        return None
    return valor if valor >= 0 else None


# ---------- Leitura dos arquivos -> lista de linhas (listas de textos) ----------

def dividir_linha(texto):
    """Quebra uma linha de texto solto em colunas (tab, ';', '|' ou 2+ espaços)."""
    partes = re.split(r"\t|;|\||\s{2,}", texto.strip())
    return [p.strip() for p in partes if p.strip()]


def ler_xlsx(dados):
    from openpyxl import load_workbook
    livro = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    linhas = []
    for planilha in livro.worksheets:
        for linha in planilha.iter_rows(values_only=True):
            celulas = ["" if v is None else (str(v) if not isinstance(v, float) else repr(v)) for v in linha]
            while celulas and not celulas[-1].strip():
                celulas.pop()
            linhas.append(celulas)
    livro.close()
    return linhas


def ler_csv(dados):
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = dados.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    amostra = texto[:4096]
    delimitador = max([";", ",", "\t", "|"], key=amostra.count)
    return [linha for linha in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def ler_docx(dados):
    from docx import Document
    documento = Document(io.BytesIO(dados))
    linhas = []
    for tabela in documento.tables:
        for linha in tabela.rows:
            celulas = []
            for celula in linha.cells:
                texto = celula.text.strip()
                # Células mescladas aparecem repetidas
                if not celulas or celulas[-1] != texto:
                    celulas.append(texto)
            linhas.append(celulas)
    if linhas:
        # Com tabela, os parágrafos costumam ser só títulos e observações
        return linhas
    for paragrafo in documento.paragraphs:
        if paragrafo.text.strip():
            linhas.append(dividir_linha(paragrafo.text))
    return linhas


def ler_pdf(dados):
    from pypdf import PdfReader
    leitor = PdfReader(io.BytesIO(dados))
    linhas = []
    for pagina in leitor.pages:
        texto = pagina.extract_text(extraction_mode="layout") or ""
        for linha in texto.splitlines():
            if linha.strip():
                linhas.append(dividir_linha(linha))
    return linhas


def ler_texto(texto):
    """Texto colado (ex.: mensagem do WhatsApp): uma linha por item."""
    linhas = []
    for linha in texto.splitlines():
        # Tira marcadores de lista, emojis e o "Nome: " das mensagens encaminhadas
        linha = re.sub(r"^\s*(\[[^\]]*\]\s*[^:]*:\s*)?", "", linha)
        linha = re.sub(r"^[\s\-–•*·>✅❌🔹🔸▪️]+", "", linha).strip()
        if linha:
            linhas.append(dividir_linha(linha))
    return linhas


def ler_arquivo(nome_arquivo, dados):
    ext = nome_arquivo.rsplit(".", 1)[-1].lower() if "." in nome_arquivo else ""
    if ext in ("xlsx", "xlsm"):
        return ler_xlsx(dados)
    if ext in ("csv", "txt"):
        return ler_csv(dados)
    if ext == "docx":
        return ler_docx(dados)
    if ext == "pdf":
        return ler_pdf(dados)
    if ext in ("doc", "xls"):
        raise ValueError(f"Arquivos .{ext} antigos não são suportados. Abra no Word/Excel e salve como "
                         f".{'docx' if ext == 'doc' else 'xlsx'}.")
    raise ValueError("Formato não suportado. Use Excel (.xlsx), CSV, Word (.docx) ou PDF.")


# ---------- Interpretação das linhas ----------

# Cabeçalho que diz que a coluna é a contagem feita ganha de "Estoque"/"Saldo" (o estoque do sistema).
CABECALHO_CONTAGEM = {"contada", "contado", "contagem", "fisico", "peso", "conferido", "conferida"}


def achar_cabecalho(linhas):
    """Procura nas primeiras linhas um cabeçalho. Retorna (índice, {campo: coluna}).
    As colunas que podem ser a quantidade ficam em mapa["_quantidades"], a preferida primeiro."""
    for i, linha in enumerate(linhas[:30]):
        mapa, quantidades = {}, []
        for col, celula in enumerate(linha):
            palavras = re.findall(r"[a-z]+", sem_acento(celula))
            for campo, chaves in CABECALHOS.items():
                if any(p in chaves for p in palavras):
                    if campo == "quantidade":
                        contagem = any(p in CABECALHO_CONTAGEM for p in palavras) and "ultima" not in palavras
                        quantidades.insert(len([c for c in quantidades if c[1]]) if contagem else len(quantidades),
                                           (col, contagem))
                    if campo not in mapa:
                        mapa[campo] = col
                        break
        if "nome" in mapa and len(mapa) >= 2:
            if quantidades:
                mapa["quantidade"] = quantidades[0][0]
                mapa["_quantidades"] = [c for c, _ in quantidades]
            return i, mapa
    return None, {}


def escolher_coluna_quantidade(corpo, mapa):
    """A coluna da quantidade precisa ter números. Se a escolhida pelo cabeçalho está toda zerada ou sem números
    (ex.: "Estoque" do sistema = 0 e a contagem ao lado), usa outra coluna depois do nome que tenha valores."""
    if "nome" not in mapa:
        return

    def contar(col):
        numeros = nao_zero = 0
        for linha in corpo:
            if col < len(linha) and str(linha[col]).strip():
                valor = parse_numero(linha[col])
                if valor is not None:
                    numeros += 1
                    nao_zero += valor > 0
        return numeros, nao_zero

    atual = mapa.get("quantidade")
    if atual is not None and contar(atual)[1] > 0:
        return
    proibidas = {mapa.get(c) for c in ("nome", "categoria", "unidade", "observacao")}
    candidatas = [c for c in mapa.get("_quantidades", []) if c != atual] + \
                 list(range(mapa["nome"] + 1, max((len(l) for l in corpo), default=0)))
    melhor, melhor_nao_zero = None, 0
    for col in candidatas:
        if col in proibidas or col == atual:
            continue
        numeros, nao_zero = contar(col)
        if nao_zero > melhor_nao_zero and numeros * 2 >= len([l for l in corpo if col < len(l) and str(l[col]).strip()]):
            melhor, melhor_nao_zero = col, nao_zero
    if melhor is not None:
        mapa["quantidade"] = melhor


def achar_categoria(texto, categorias):
    """Casa o texto com uma categoria existente pela primeira palavra significativa."""
    alvo = re.findall(r"[a-z0-9]+", sem_acento(texto))
    if not alvo:
        return None
    for categoria in categorias:
        nome = re.findall(r"[a-z0-9]+", sem_acento(categoria))
        if alvo == nome or (nome and len(nome[0]) >= 4 and alvo[0] == nome[0]):
            return categoria
    return None


def interpretar_sem_cabecalho(celulas, categorias=()):
    """Sem cabeçalho: o primeiro texto é o nome; um número é a quantidade; uma unidade é a unidade.
    Planilha "código; grupo; nome; un; 0; 3,3 KG": o código antes do nome não é quantidade, o grupo vira a
    categoria, e entre vários números depois do nome vale o que tem a unidade escrita junto ou o que não é zero."""
    linha = {"nome": "", "unidade": None, "quantidade": None, "categoria": ""}
    textos = [c for c in celulas if c.strip() and parse_numero(c) is None and normalizar_unidade(c) is None]
    numeros_antes, depois = [], []  # (valor, unidade escrita junto)
    for celula in celulas:
        if not celula.strip():
            continue
        numero = parse_numero(celula)
        unidade = normalizar_unidade(celula)
        if not linha["nome"] and numero is None and unidade is None:
            if len(textos) > 1 and not linha["categoria"] and achar_categoria(celula, categorias):
                linha["categoria"] = celula.strip()
                continue
            linha["nome"] = celula.strip()
        elif unidade and not linha["unidade"]:
            linha["unidade"] = unidade
        elif numero is not None:
            (depois if linha["nome"] else numeros_antes).append((numero, separar_quantidade(celula)[1]))
        else:
            # "2 kg" junto na mesma célula
            m = re.fullmatch(r"([\d.,]+)\s*([^\d\s].*)", celula.strip())
            if m and parse_numero(m.group(1)) is not None and normalizar_unidade(m.group(2)):
                depois.append((parse_numero(m.group(1)), normalizar_unidade(m.group(2))))
    numeros = depois or numeros_antes
    if numeros:
        escolhido = next((n for n in numeros if n[1]), None) or next((n for n in numeros if n[0] > 0), numeros[0])
        linha["quantidade"] = escolhido[0]
        linha["unidade"] = escolhido[1] or linha["unidade"]
        linha["unidade_na_qtd"] = bool(escolhido[1])
    # "Lentilha 2 kg", "Lentilha: 2kg", "Lentilha - 2" numa célula só
    m = re.fullmatch(r"(.+?)[\s:=\-]+([\d.,]+)\s*([^\d\s.]*)\.?", linha["nome"])
    if m and parse_numero(m.group(2)) is not None and (not m.group(3) or normalizar_unidade(m.group(3))):
        linha["nome"] = m.group(1).strip(" :-")
        linha["quantidade"] = linha["quantidade"] if linha["quantidade"] is not None else parse_numero(m.group(2))
        if m.group(3):
            linha["unidade"] = linha["unidade"] or normalizar_unidade(m.group(3))
            linha["unidade_na_qtd"] = True
    # "2 kg de lentilha" / "2kg lentilha"
    m = re.fullmatch(r"([\d.,]+)\s*([^\d\s]+)\s+(?:de\s+)?(.+)", linha["nome"])
    if m and parse_numero(m.group(1)) is not None and normalizar_unidade(m.group(2)):
        linha["nome"] = m.group(3).strip()
        linha["quantidade"] = parse_numero(m.group(1))
        linha["unidade"] = normalizar_unidade(m.group(2))
        linha["unidade_na_qtd"] = True
    return linha


def extrair_itens(linhas, categorias):
    """Transforma as linhas do arquivo em itens {nome, categoria, unidade, quantidade}."""
    inicio, mapa = achar_cabecalho(linhas)
    corpo = linhas[inicio + 1:] if inicio is not None else linhas
    escolher_coluna_quantidade(corpo, mapa)
    itens, vistos = [], set()
    categoria_atual = ""

    def celula(linha, campo):
        col = mapa.get(campo)
        return linha[col].strip() if col is not None and col < len(linha) and linha[col] else ""

    for linha in corpo:
        preenchidas = [c for c in linha if str(c).strip()]
        if not preenchidas:
            continue
        # Linha com um texto só que é nome de categoria (ex.: "🥩 Carnes, Embutidos e Pescados")
        if len(preenchidas) == 1:
            categoria = achar_categoria(preenchidas[0], categorias)
            if categoria:
                categoria_atual = categoria
                continue

        if mapa:
            item = {
                "nome": celula(linha, "nome"),
                "categoria": celula(linha, "categoria"),
                # "3,3 KG" escrito na quantidade vale mais que a coluna da unidade
                "unidade_na_qtd": bool(separar_quantidade(celula(linha, "quantidade"))[1]),
                "unidade": separar_quantidade(celula(linha, "quantidade"))[1] or normalizar_unidade(celula(linha, "unidade")),
                "quantidade": parse_numero(celula(linha, "quantidade")),
                "observacao": celula(linha, "observacao"),
            }
            # PDF perde as células vazias (ex.: Local em branco) e as seguintes andam para a esquerda:
            # se a coluna da quantidade não tem número, usa o primeiro número depois do nome.
            depois = [str(c) for c in linha[mapa["nome"] + 1:]] if "nome" in mapa else []
            if item["quantidade"] is None:
                item["quantidade"] = next((parse_numero(c) for c in depois if parse_numero(c) is not None), None)
            if not item["unidade"]:
                item["unidade"] = next((normalizar_unidade(c) or separar_quantidade(c)[1] for c in depois
                                        if normalizar_unidade(c) or separar_quantidade(c)[1]), None)
        else:
            item = interpretar_sem_cabecalho([str(c) for c in linha], categorias)

        nome = re.sub(r"\s+", " ", item["nome"]).strip(" -:•*")
        nome, unidade_no_nome = tirar_unidade_do_nome(nome)
        if unidade_no_nome and not item.get("unidade_na_qtd"):
            # "Bife do Vazio Empanado kg": a unidade escrita junto do nome vale mais que a coluna Un
            item["unidade"], item["unidade_na_qtd"] = unidade_no_nome, True
        if not nome or parse_numero(nome) is not None or len(nome) > 80:
            continue
        # Linhas de total/rodapé
        if sem_acento(nome).split(" ")[0] in ("total", "subtotal", "pagina", "data", "responsavel"):
            continue
        item["nome"] = nome
        item["categoria"] = achar_categoria(item["categoria"], categorias) or item["categoria"] or categoria_atual
        chave = sem_acento(nome)
        if chave in vistos:
            continue
        vistos.add(chave)
        itens.append(item)
    return itens
