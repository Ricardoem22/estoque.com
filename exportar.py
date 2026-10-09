# exportar.py
# Uma tabela (título, cabeçalho e linhas) vira arquivo para baixar: CSV, Excel (.xlsx), PDF ou Word (.docx).
# Cada tela monta as linhas uma vez e escolhe o formato pelo ?formato= do link.
import csv
import io
import re

from flask import Response, request

FORMATOS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
}


def formato_pedido():
    formato = request.args.get("formato", "csv")
    return formato if formato in FORMATOS else "csv"


def _csv(titulo, subtitulo, cabecalho, linhas):
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";")
    if subtitulo:
        escritor.writerow([subtitulo])
        escritor.writerow([])
    escritor.writerow(cabecalho)
    escritor.writerows(linhas)
    # BOM para o Excel reconhecer os acentos
    return ("﻿" + saida.getvalue()).encode("utf-8")


def _xlsx(titulo, subtitulo, cabecalho, linhas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    livro = Workbook()
    folha = livro.active
    folha.title = re.sub(r"[\\/*?:\[\]]", "-", titulo)[:31]
    folha.append([titulo])
    folha["A1"].font = Font(bold=True, size=14)
    if subtitulo:
        folha.append([subtitulo])
    folha.append([])
    folha.append(cabecalho)
    linha_cab = folha.max_row
    for celula in folha[linha_cab]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="FFE9A8")
    for linha in linhas:
        folha.append([_numero(v) for v in linha])
    for i, nome in enumerate(cabecalho, 1):
        largura = max([len(str(nome))] + [len(str(l[i - 1])) for l in linhas if i - 1 < len(l)])
        folha.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 8), 45)
    folha.freeze_panes = folha.cell(row=linha_cab + 1, column=1)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _numero(valor):
    """No Excel, "3,5" vira número de verdade para dar para somar."""
    if isinstance(valor, str):
        texto = valor.strip()
        if texto and texto.replace(",", "", 1).replace("-", "", 1).isdigit() and texto.count(",") <= 1:
            return float(texto.replace(",", "."))
    return valor


def _latin(texto):
    # As fontes padrão do PDF não têm emoji; acentos do português ficam
    return str(texto).encode("latin-1", "ignore").decode("latin-1").strip()


def _pdf(titulo, subtitulo, cabecalho, linhas):
    from fpdf import FPDF
    pdf = FPDF(orientation="L" if len(cabecalho) > 5 else "P", format="A4")
    pdf.set_auto_page_break(True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 9, _latin(titulo), new_x="LMARGIN", new_y="NEXT")
    if subtitulo:
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 6, _latin(subtitulo), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Helvetica", "", 9)
    from fpdf.fonts import FontFace
    cabeca = FontFace(emphasis="BOLD", fill_color=(255, 233, 168))
    with pdf.table(text_align="LEFT", line_height=5, headings_style=cabeca) as tabela:
        for linha in [cabecalho] + [list(l) for l in linhas]:
            fila = tabela.row()
            for valor in linha:
                fila.cell(_latin(valor if valor is not None else ""))
    return bytes(pdf.output())


def _docx(titulo, subtitulo, cabecalho, linhas):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt
    doc = Document()
    if len(cabecalho) > 5:
        secao = doc.sections[0]
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = secao.page_height, secao.page_width
    doc.add_heading(titulo, level=1)
    if subtitulo:
        doc.add_paragraph(subtitulo)
    tabela = doc.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Light Grid Accent 1"
    for celula, nome in zip(tabela.rows[0].cells, cabecalho):
        celula.text = str(nome)
    for linha in linhas:
        for celula, valor in zip(tabela.add_row().cells, linha):
            celula.text = "" if valor is None else str(valor)
    for fila in tabela.rows:
        for celula in fila.cells:
            for paragrafo in celula.paragraphs:
                for trecho in paragrafo.runs:
                    trecho.font.size = Pt(9)
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def responder(nome_arquivo, titulo, cabecalho, linhas, subtitulo="", formato=None):
    """Devolve o arquivo para baixar. nome_arquivo sem extensão."""
    formato = formato or formato_pedido()
    gerar = {"csv": _csv, "xlsx": _xlsx, "pdf": _pdf, "docx": _docx}[formato]
    mimetype, extensao = FORMATOS[formato]
    return Response(gerar(titulo, subtitulo, cabecalho, linhas), mimetype=mimetype,
                    headers={"Content-Disposition": f"attachment; filename={nome_arquivo}.{extensao}"})
